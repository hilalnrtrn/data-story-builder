from __future__ import annotations

import random
import copy
import csv
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook


# ============================================================
# PROJECT PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

SAMPLE_DATA_DIR = PROJECT_ROOT / "sample-data"

CLEAN_CSV_PATH = SAMPLE_DATA_DIR / "sales_demo_clean.csv"
DIRTY_CSV_PATH = SAMPLE_DATA_DIR / "sales_demo_dirty.csv"
EXCEL_PATH = SAMPLE_DATA_DIR / "sales_demo.xlsx"


# ============================================================
# REPRODUCIBILITY
# ============================================================

RANDOM_SEED = 42


def initialize_random_seed() -> None:
    """
    Ensures that controlled pseudo-random values are reproducible.

    Running the generator multiple times with the same seed should
    produce the same synthetic dataset.
    """
    random.seed(RANDOM_SEED)


# ============================================================
# DATASET DIMENSIONS
# ============================================================

REGIONS = [
    "Marmara",
    "Ege",
    "İç Anadolu",
    "Akdeniz",
    "Karadeniz",
]

PRODUCT_CATEGORY_MAP = {
    "Wireless Headphones": "Electronics",
    "Smart Watch": "Electronics",
    "Bluetooth Speaker": "Electronics",
    "Power Bank": "Electronics",

    "Coffee Machine": "Home",
    "Air Fryer": "Home",
    "Vacuum Cleaner": "Home",
    "Desk Lamp": "Home",

    "Office Chair": "Office",
    "Mechanical Keyboard": "Office",
    "Monitor Stand": "Office",
    "Webcam": "Office",
}

CHANNELS = [
    "Online",
    "Store",
]

CAMPAIGNS = [
    "None",
    "Summer Sale",
    "Back to School",
    "Black Friday",
    "Holiday Campaign",
]
def get_campaign_for_date(date_string: str) -> str:
    """
    Returns the campaign context for a given date.
    """

    month = date_string[:7]

    if month == "2025-06":
        return "Summer Sale"

    if month in {"2025-08", "2025-09"}:
        return "Back to School"

    if month == "2025-11":
        return "Black Friday"

    if month == "2025-12":
        return "Holiday Campaign"

    return "None"


# ============================================================
# CONTROLLED MONTHLY TARGETS
# ============================================================

MONTHLY_REVENUE_TARGETS = {
    "2025-01": 900_000,
    "2025-02": 920_000,
    "2025-03": 950_000,
    "2025-04": 980_000,
    "2025-05": 1_000_000,
    "2025-06": 1_300_000,
    "2025-07": 1_040_000,
    "2025-08": 1_060_000,
    "2025-09": 1_080_000,
    "2025-10": 1_100_000,
    "2025-11": 1_130_000,
    "2025-12": 1_160_000,
}

MONTHLY_ORDER_TARGETS = {
    "2025-01": 8_900,
    "2025-02": 9_100,
    "2025-03": 9_350,
    "2025-04": 9_700,
    "2025-05": 10_000,
    "2025-06": 8_500,
    "2025-07": 9_750,
    "2025-08": 9_900,
    "2025-09": 10_050,
    "2025-10": 10_200,
    "2025-11": 10_350,
    "2025-12": 10_500,
}

MONTHLY_CUSTOMER_TARGETS = {
    "2025-01": 7_200,
    "2025-02": 7_350,
    "2025-03": 7_500,
    "2025-04": 7_750,
    "2025-05": 8_000,
    "2025-06": 7_040,
    "2025-07": 7_800,
    "2025-08": 7_900,
    "2025-09": 8_000,
    "2025-10": 8_100,
    "2025-11": 8_200,
    "2025-12": 8_300,
}

MONTHLY_MARKETING_SPEND_TARGETS = {
    "2025-01": 82_000,
    "2025-02": 84_000,
    "2025-03": 88_000,
    "2025-04": 94_000,
    "2025-05": 100_000,
    "2025-06": 120_000,
    "2025-07": 98_000,
    "2025-08": 102_000,
    "2025-09": 108_000,
    "2025-10": 110_000,
    "2025-11": 118_000,
    "2025-12": 125_000,
}

MONTHLY_QUANTITY_TARGETS = {
    "2025-01": 11_200,
    "2025-02": 11_450,
    "2025-03": 11_800,
    "2025-04": 12_100,
    "2025-05": 12_500,
    "2025-06": 11_600,
    "2025-07": 12_200,
    "2025-08": 12_400,
    "2025-09": 12_600,
    "2025-10": 12_850,
    "2025-11": 13_100,
    "2025-12": 13_300,
}

MONTHLY_DISCOUNT_TARGETS = {
    "2025-01": 0.05,
    "2025-02": 0.05,
    "2025-03": 0.06,
    "2025-04": 0.06,
    "2025-05": 0.08,
    "2025-06": 0.12,
    "2025-07": 0.09,
    "2025-08": 0.10,
    "2025-09": 0.18,
    "2025-10": 0.11,
    "2025-11": 0.15,
    "2025-12": 0.16,
}

MONTHLY_RETURN_RATE_TARGETS = {
    "2025-01": 0.020,
    "2025-02": 0.021,
    "2025-03": 0.020,
    "2025-04": 0.021,
    "2025-05": 0.022,
    "2025-06": 0.023,
    "2025-07": 0.021,
    "2025-08": 0.020,
    "2025-09": 0.035,
    "2025-10": 0.023,
    "2025-11": 0.024,
    "2025-12": 0.025,
}


# ============================================================
# CONTROLLED REVENUE BREAKDOWNS
# ============================================================

REGION_REVENUE_TARGETS = {
    "2025-05": {
        "Marmara": 380_000,
        "Ege": 220_000,
        "Akdeniz": 160_000,
        "İç Anadolu": 150_000,
        "Karadeniz": 90_000,
    },
    "2025-06": {
        "Marmara": 590_000,
        "Ege": 260_000,
        "Akdeniz": 190_000,
        "İç Anadolu": 135_000,
        "Karadeniz": 125_000,
    },
}

CHANNEL_REVENUE_TARGETS = {
    "2025-05": {
        "Online": 550_000,
        "Store": 450_000,
    },
    "2025-06": {
        "Online": 780_000,
        "Store": 520_000,
    },
}


# ============================================================
# CONTROLLED PRODUCT EVENTS
# ============================================================

SMART_WATCH_MONTHLY_REVENUE = {
    "2025-01": 70_000,
    "2025-02": 72_000,
    "2025-03": 71_000,
    "2025-04": 74_000,
    "2025-05": 73_000,
    "2025-06": 145_000,
    "2025-07": 75_000,
    "2025-08": 76_000,
    "2025-09": 74_000,
    "2025-10": 77_000,
    "2025-11": 78_000,
    "2025-12": 80_000,
}

OFFICE_CHAIR_MONTHLY_REVENUE = {
    "2025-01": 45_000,
    "2025-02": 49_000,
    "2025-03": 54_000,
    "2025-04": 58_000,
    "2025-05": 63_000,
    "2025-06": 69_000,
    "2025-07": 74_000,
    "2025-08": 80_000,
    "2025-09": 86_000,
    "2025-10": 91_000,
    "2025-11": 97_000,
    "2025-12": 103_000,
}

BLUETOOTH_SPEAKER_MONTHLY_REVENUE = {
    "2025-01": 50_000,
    "2025-02": 51_000,
    "2025-03": 50_000,
    "2025-04": 52_000,
    "2025-05": 51_000,
    "2025-06": 52_000,
    "2025-07": 51_000,
    "2025-08": 53_000,
    "2025-09": 140_000,
    "2025-10": 52_000,
    "2025-11": 53_000,
    "2025-12": 51_000,
}
# ============================================================
# DEFAULT ALLOCATION WEIGHTS
# ============================================================

DEFAULT_REGION_SHARES = {
    "Marmara": 0.38,
    "Ege": 0.22,
    "Akdeniz": 0.16,
    "İç Anadolu": 0.15,
    "Karadeniz": 0.09,
}

DEFAULT_CHANNEL_SHARES = {
    "Online": 0.55,
    "Store": 0.45,
}


OTHER_PRODUCT_WEIGHTS = {
    "Wireless Headphones": 1.15,
    "Power Bank": 0.95,
    "Coffee Machine": 1.10,
    "Air Fryer": 1.05,
    "Vacuum Cleaner": 1.20,
    "Desk Lamp": 0.75,
    "Mechanical Keyboard": 0.95,
    "Monitor Stand": 0.80,
    "Webcam": 0.85,
}

# ============================================================
# DATASET SCHEMA
# ============================================================

EXPECTED_COLUMNS = [
    "Date",
    "Region",
    "Category",
    "Product",
    "Channel",
    "Campaign",
    "Revenue",
    "Orders",
    "Customers",
    "Quantity",
    "MarketingSpend",
    "StockLevel",
    "Discount",
    "Returns",
]

# ============================================================
# CONTROLLED DIRTY DATASET ISSUES
# ============================================================

DIRTY_ISSUE_COUNTS = {
    "duplicate_rows": 10,
    "region_normalization_issues": 12,
    "product_whitespace_issues": 5,
    "invalid_product_category_mappings": 3,
    "stock_level_missing_values": 50,
    "invalid_discount_values": 2,
    "negative_order_values": 2,
    "returns_greater_than_quantity": 2,
    "invalid_date_values": 2,
    "numeric_format_issues": 4,
}


@dataclass(frozen=True)
class DatasetConfig:
    start_date: str = "2025-01-01"
    end_date: str = "2025-12-31"
    frequency: str = "DAILY"
    random_seed: int = RANDOM_SEED


CONFIG = DatasetConfig()

# ============================================================
# REVENUE ALLOCATION
# ============================================================

def build_product_revenue_targets(month: str) -> dict[str, float]:
    """
    Creates monthly Revenue targets for all 12 products.

    Three products have explicitly controlled Golden Scenario values:
    - Smart Watch
    - Office Chair
    - Bluetooth Speaker

    Remaining Revenue is distributed across the other products using
    deterministic weights.
    """

    monthly_total = MONTHLY_REVENUE_TARGETS[month]

    controlled_targets = {
        "Smart Watch": SMART_WATCH_MONTHLY_REVENUE[month],
        "Office Chair": OFFICE_CHAIR_MONTHLY_REVENUE[month],
        "Bluetooth Speaker": BLUETOOTH_SPEAKER_MONTHLY_REVENUE[month],
    }

    controlled_total = sum(controlled_targets.values())

    remaining_revenue = monthly_total - controlled_total

    assert remaining_revenue > 0, (
        f"Controlled product targets exceed monthly Revenue for {month}"
    )

    weight_total = sum(OTHER_PRODUCT_WEIGHTS.values())

    product_targets = dict(controlled_targets)

    for product, weight in OTHER_PRODUCT_WEIGHTS.items():
        share = weight / weight_total
        product_targets[product] = remaining_revenue * share

    assert len(product_targets) == 12, (
        f"Expected 12 product Revenue targets for {month}"
    )

    return product_targets


def get_region_shares(month: str) -> dict[str, float]:
    """
    Returns the Region Revenue shares for a month.

    May and June use explicitly controlled Golden Scenario targets.
    Other months use stable default synthetic shares.
    """

    if month in REGION_REVENUE_TARGETS:
        monthly_total = MONTHLY_REVENUE_TARGETS[month]

        return {
            region: revenue / monthly_total
            for region, revenue in REGION_REVENUE_TARGETS[month].items()
        }

    return DEFAULT_REGION_SHARES.copy()


def get_channel_shares(month: str) -> dict[str, float]:
    """
    Returns Channel Revenue shares for a month.

    May and June use explicitly controlled targets.
    Other months use default synthetic shares.
    """

    if month in CHANNEL_REVENUE_TARGETS:
        monthly_total = MONTHLY_REVENUE_TARGETS[month]

        return {
            channel: revenue / monthly_total
            for channel, revenue in CHANNEL_REVENUE_TARGETS[month].items()
        }

    return DEFAULT_CHANNEL_SHARES.copy()


def build_monthly_revenue_cube(
    month: str,
) -> dict[tuple[str, str, str], float]:
    """
    Allocates monthly Revenue across:

        Region × Product × Channel

    The allocation preserves the required monthly marginal totals.

    Key:
        (Region, Product, Channel)

    Value:
        Monthly Revenue for that combination.
    """

    product_targets = build_product_revenue_targets(month)
    region_shares = get_region_shares(month)
    channel_shares = get_channel_shares(month)

    cube: dict[tuple[str, str, str], float] = {}

    for region in REGIONS:
        for product in PRODUCT_CATEGORY_MAP:
            for channel in CHANNELS:
                revenue = (
                    product_targets[product]
                    * region_shares[region]
                    * channel_shares[channel]
                )

                cube[(region, product, channel)] = revenue

    return cube
def validate_monthly_revenue_cube(
    month: str,
    cube: dict[tuple[str, str, str], float],
) -> None:
    """
    Validates that monthly Revenue allocation preserves all required
    controlled totals.
    """

    tolerance = 0.01
    expected_total = MONTHLY_REVENUE_TARGETS[month]

    # --------------------------------------------------------
    # TOTAL REVENUE
    # --------------------------------------------------------

    actual_total = sum(cube.values())

    assert abs(actual_total - expected_total) <= tolerance, (
        f"Total Revenue mismatch for {month}: "
        f"{actual_total} != {expected_total}"
    )

    # --------------------------------------------------------
    # PRODUCT TOTALS
    # --------------------------------------------------------

    expected_product_targets = build_product_revenue_targets(month)

    for product, expected_product_total in expected_product_targets.items():

        actual_product_total = sum(
            revenue
            for (region, cube_product, channel), revenue in cube.items()
            if cube_product == product
        )

        assert abs(actual_product_total - expected_product_total) <= tolerance, (
            f"Product Revenue mismatch for {month} / {product}: "
            f"{actual_product_total} != {expected_product_total}"
        )

    # --------------------------------------------------------
    # REGION TOTALS
    # --------------------------------------------------------

    if month in REGION_REVENUE_TARGETS:

        for region, expected_region_total in REGION_REVENUE_TARGETS[month].items():

            actual_region_total = sum(
                revenue
                for (cube_region, product, channel), revenue in cube.items()
                if cube_region == region
            )

            assert abs(actual_region_total - expected_region_total) <= tolerance, (
                f"Region Revenue mismatch for {month} / {region}: "
                f"{actual_region_total} != {expected_region_total}"
            )

    # --------------------------------------------------------
    # CHANNEL TOTALS
    # --------------------------------------------------------

    if month in CHANNEL_REVENUE_TARGETS:

        for channel, expected_channel_total in CHANNEL_REVENUE_TARGETS[month].items():

            actual_channel_total = sum(
                revenue
                for (region, product, cube_channel), revenue in cube.items()
                if cube_channel == channel
            )

            assert abs(actual_channel_total - expected_channel_total) <= tolerance, (
                f"Channel Revenue mismatch for {month} / {channel}: "
                f"{actual_channel_total} != {expected_channel_total}"
            )
def build_and_validate_revenue_cubes(
) -> dict[str, dict[tuple[str, str, str], float]]:
    """
    Builds and validates Revenue allocation for every month.
    """

    revenue_cubes = {}

    for month in MONTHLY_REVENUE_TARGETS:
        cube = build_monthly_revenue_cube(month)

        validate_monthly_revenue_cube(
            month=month,
            cube=cube,
        )

        revenue_cubes[month] = cube

    return revenue_cubes

# ============================================================
# DAILY REVENUE DISTRIBUTION
# ============================================================

def get_dates_for_month(month: str) -> list[date]:
    """
    Returns all calendar dates belonging to a YYYY-MM month.
    """

    year, month_number = map(int, month.split("-"))
    days_in_month = monthrange(year, month_number)[1]

    return [
        date(year, month_number, day)
        for day in range(1, days_in_month + 1)
    ]


def build_daily_weights(day_count: int) -> list[float]:
    """
    Creates deterministic daily variation while ensuring that
    all weights sum exactly to 1.0.

    The random seed initialized at startup keeps the generated
    pattern reproducible.
    """

    raw_weights = [
        random.uniform(0.85, 1.15)
        for _ in range(day_count)
    ]

    total_weight = sum(raw_weights)

    return [
        weight / total_weight
        for weight in raw_weights
    ]


def distribute_monthly_revenue_to_daily_rows(
    month: str,
    cube: dict[tuple[str, str, str], float],
) -> list[dict]:
    """
    Converts one monthly Region × Product × Channel Revenue cube
    into daily grain rows.

    Final grain:

        Date × Region × Product × Channel
    """

    dates = get_dates_for_month(month)

    rows: list[dict] = []

    for (region, product, channel), monthly_revenue in cube.items():

        daily_weights = build_daily_weights(len(dates))

        for current_date, daily_weight in zip(dates, daily_weights):

            daily_revenue = monthly_revenue * daily_weight

            rows.append(
                {
                    "Date": current_date.isoformat(),
                    "Region": region,
                    "Category": PRODUCT_CATEGORY_MAP[product],
                    "Product": product,
                    "Channel": channel,
                    "Revenue": daily_revenue,
                }
            )

    return rows


def build_daily_revenue_rows(
    revenue_cubes: dict[str, dict[tuple[str, str, str], float]],
) -> list[dict]:
    """
    Builds the full-year daily Revenue dataset skeleton.
    """

    all_rows: list[dict] = []

    for month, cube in revenue_cubes.items():
        month_rows = distribute_monthly_revenue_to_daily_rows(
            month=month,
            cube=cube,
        )

        all_rows.extend(month_rows)

    return all_rows

# ============================================================
# ADDITIONAL METRIC ALLOCATION
# ============================================================

def allocate_monthly_metric_to_rows(
    rows: list[dict],
    month: str,
    metric_name: str,
    monthly_target: float,
) -> None:
    """
    Allocates a monthly metric target across all rows belonging
    to the given month.

    Revenue is used as the allocation basis so that larger
    commercial observations generally receive larger shares.

    The final monthly total is preserved.
    """

    month_rows = [
        row
        for row in rows
        if row["Date"].startswith(month)
    ]

    assert month_rows, f"No rows found for month: {month}"

    total_revenue = sum(
        row["Revenue"]
        for row in month_rows
    )

    assert total_revenue > 0, (
        f"Revenue must be positive for metric allocation: {month}"
    )

    allocated_total = 0.0

    for index, row in enumerate(month_rows):

        if index == len(month_rows) - 1:
            value = monthly_target - allocated_total
        else:
            revenue_share = row["Revenue"] / total_revenue
            value = monthly_target * revenue_share
            allocated_total += value

        row[metric_name] = value

def allocate_monthly_integer_metric_to_rows(
    rows: list[dict],
    month: str,
    metric_name: str,
    monthly_target: int,
    weight_metric: str = "Revenue",
) -> None:
    """
    Allocates an integer monthly target across the month's rows.

    Uses the largest-remainder method so that:
    - every row receives an integer value,
    - the exact monthly target is preserved,
    - allocation remains proportional to the selected weight metric.
    """

    month_rows = [
        row
        for row in rows
        if row["Date"].startswith(month)
    ]

    assert month_rows, f"No rows found for month: {month}"

    total_weight = sum(
        row[weight_metric]
        for row in month_rows
    )

    assert total_weight > 0, (
        f"{weight_metric} must be positive for {month}"
    )

    allocations = []

    for row in month_rows:
        exact_value = (
            monthly_target
            * row[weight_metric]
            / total_weight
        )

        base_value = int(exact_value)

        allocations.append(
            {
                "row": row,
                "base": base_value,
                "remainder": exact_value - base_value,
            }
        )

    allocated_total = sum(
        item["base"]
        for item in allocations
    )

    remaining_units = monthly_target - allocated_total

    allocations.sort(
        key=lambda item: item["remainder"],
        reverse=True,
    )

    for index, item in enumerate(allocations):
        value = item["base"]

        if index < remaining_units:
            value += 1

        item["row"][metric_name] = value

def add_controlled_metrics(
    rows: list[dict],
) -> None:
    """
    Adds Orders, Customers, and MarketingSpend to the existing
    daily Revenue rows using controlled monthly targets.
    """

    for month in MONTHLY_REVENUE_TARGETS:

        allocate_monthly_integer_metric_to_rows(
            rows=rows,
            month=month,
            metric_name="Orders",
            monthly_target=MONTHLY_ORDER_TARGETS[month],
        )

        allocate_monthly_integer_metric_to_rows(
            rows=rows,
            month=month,
            metric_name="Customers",
            monthly_target=MONTHLY_CUSTOMER_TARGETS[month],
        )

        allocate_monthly_metric_to_rows(
            rows=rows,
            month=month,
            metric_name="MarketingSpend",
            monthly_target=MONTHLY_MARKETING_SPEND_TARGETS[month],
        )

def add_quantity_metric(
    rows: list[dict],
) -> None:
    """
    Adds controlled integer Quantity values using monthly targets.
    """

    for month in MONTHLY_QUANTITY_TARGETS:
        allocate_monthly_integer_metric_to_rows(
            rows=rows,
            month=month,
            metric_name="Quantity",
            monthly_target=MONTHLY_QUANTITY_TARGETS[month],
        )


def add_discount_metric(
    rows: list[dict],
) -> None:
    """
    Adds monthly controlled average Discount values.
    """

    for row in rows:
        month = row["Date"][:7]
        row["Discount"] = MONTHLY_DISCOUNT_TARGETS[month]


def add_returns_metric(
    rows: list[dict],
) -> None:
    """
    Creates integer Returns values while preserving the controlled
    monthly ReturnRate as closely and deterministically as possible.
    """

    for month, return_rate in MONTHLY_RETURN_RATE_TARGETS.items():

        monthly_quantity = sum(
            row["Quantity"]
            for row in rows
            if row["Date"].startswith(month)
        )

        target_returns = round(
            monthly_quantity * return_rate
        )

        allocate_monthly_integer_metric_to_rows(
            rows=rows,
            month=month,
            metric_name="Returns",
            monthly_target=target_returns,
            weight_metric="Quantity",
        )


def add_campaign_dimension(
    rows: list[dict],
) -> None:
    """
    Adds Campaign context to every row.
    """

    for row in rows:
        row["Campaign"] = get_campaign_for_date(
            row["Date"]
        )

def add_stock_level_metric(
    rows: list[dict],
) -> None:
    """
    Adds deterministic synthetic StockLevel values.

    StockLevel is a snapshot metric and is therefore generated
    independently rather than allocated as a monthly additive total.
    """

    for row in rows:

        product_factor = (
            list(PRODUCT_CATEGORY_MAP.keys()).index(
                row["Product"]
            )
            + 1
        )

        region_factor = (
            REGIONS.index(row["Region"])
            + 1
        )

        channel_factor = (
            CHANNELS.index(row["Channel"])
            + 1
        )

        base_stock = (
            120
            + product_factor * 8
            + region_factor * 5
            + channel_factor * 3
        )

        variation = random.randint(-15, 15)

        row["StockLevel"] = max(
            0,
            base_stock + variation,
        )

def add_remaining_dataset_fields(
    rows: list[dict],
) -> None:
    """
    Adds the remaining clean-dataset fields.
    """

    add_quantity_metric(rows)
    add_discount_metric(rows)
    add_returns_metric(rows)
    add_stock_level_metric(rows)
    add_campaign_dimension(rows)

def validate_remaining_dataset_fields(
    rows: list[dict],
) -> None:
    """
    Validates clean dataset rules for Quantity, Discount,
    Returns, StockLevel, and Campaign.
    """
    assert all(
        isinstance(row["Orders"], int)
        for row in rows
    ), "Orders must contain integer values."

    assert all(
        isinstance(row["Customers"], int)
        for row in rows
    ), "Customers must contain integer values."

    assert all(
        isinstance(row["Quantity"], int)
        for row in rows
    ), "Quantity must contain integer values."

    assert all(
        isinstance(row["Returns"], int)
        for row in rows
    ), "Returns must contain integer values."

    assert all(
        isinstance(row["StockLevel"], int)
        for row in rows
    ), "StockLevel must contain integer values."

    assert all(
        row["Quantity"] >= 0
        for row in rows
    ), "Quantity must not be negative."

    assert all(
        0.0 <= row["Discount"] <= 0.50
        for row in rows
    ), "Discount must stay between 0 and 0.50."

    assert all(
        row["Returns"] >= 0
        for row in rows
    ), "Returns must not be negative."

    assert all(
        row["Returns"] <= row["Quantity"]
        for row in rows
    ), "Returns must not exceed Quantity."

    assert all(
        row["StockLevel"] >= 0
        for row in rows
    ), "StockLevel must not be negative."

    assert all(
        row["Campaign"] in CAMPAIGNS
        for row in rows
    ), "Unexpected Campaign value detected."

# ============================================================
# FINAL CLEAN DATASET VALIDATION
# ============================================================

def validate_clean_dataset(
    rows: list[dict],
) -> None:
    """
    Performs final structural and semantic validation of the
    complete clean dataset before export.

    This function acts as the final gate before a dataset is
    considered valid enough to be written to disk.
    """

    expected_row_count = (
        365
        * len(REGIONS)
        * len(PRODUCT_CATEGORY_MAP)
        * len(CHANNELS)
    )

    # --------------------------------------------------------
    # STRUCTURE
    # --------------------------------------------------------

    assert len(rows) == expected_row_count, (
        f"Expected {expected_row_count} rows, "
        f"but found {len(rows)}."
    )

    assert rows, "Clean dataset must not be empty."

    expected_column_set = set(EXPECTED_COLUMNS)

    for index, row in enumerate(rows):
        actual_columns = set(row.keys())

        assert actual_columns == expected_column_set, (
            f"Column mismatch at row {index}. "
            f"Expected {expected_column_set}, "
            f"found {actual_columns}."
        )

    # --------------------------------------------------------
    # DATE VALIDATION
    # --------------------------------------------------------

    expected_dates = set()

    current_date = date.fromisoformat(CONFIG.start_date)
    final_date = date.fromisoformat(CONFIG.end_date)

    while current_date <= final_date:
        expected_dates.add(current_date.isoformat())
        current_date += timedelta(days=1)

    actual_dates = {
        row["Date"]
        for row in rows
    }

    assert actual_dates == expected_dates, (
        "Dataset does not contain the expected complete "
        "2025 date range."
    )

    expected_rows_per_date = (
        len(REGIONS)
        * len(PRODUCT_CATEGORY_MAP)
        * len(CHANNELS)
    )

    for expected_date in expected_dates:

        actual_count = sum(
            1
            for row in rows
            if row["Date"] == expected_date
        )

        assert actual_count == expected_rows_per_date, (
            f"Unexpected row count for {expected_date}: "
            f"{actual_count} != {expected_rows_per_date}"
        )

    # --------------------------------------------------------
    # DIMENSION VALIDATION
    # --------------------------------------------------------

    assert all(
        row["Region"] in REGIONS
        for row in rows
    ), "Invalid Region detected in clean dataset."

    assert all(
        row["Product"] in PRODUCT_CATEGORY_MAP
        for row in rows
    ), "Invalid Product detected in clean dataset."

    assert all(
        row["Category"]
        == PRODUCT_CATEGORY_MAP[row["Product"]]
        for row in rows
    ), "Product-Category mapping inconsistency detected."

    assert all(
        row["Channel"] in CHANNELS
        for row in rows
    ), "Invalid Channel detected in clean dataset."

    assert all(
        row["Campaign"] in CAMPAIGNS
        for row in rows
    ), "Invalid Campaign detected in clean dataset."

    # --------------------------------------------------------
    # NUMERIC TYPE VALIDATION
    # --------------------------------------------------------

    integer_metrics = [
        "Orders",
        "Customers",
        "Quantity",
        "StockLevel",
        "Returns",
    ]

    for metric in integer_metrics:
        assert all(
            isinstance(row[metric], int)
            and not isinstance(row[metric], bool)
            for row in rows
        ), f"{metric} must contain integer values."

    numeric_metrics = [
        "Revenue",
        "MarketingSpend",
        "Discount",
    ]

    for metric in numeric_metrics:
        assert all(
            isinstance(row[metric], (int, float))
            and not isinstance(row[metric], bool)
            for row in rows
        ), f"{metric} must contain numeric values."

    # --------------------------------------------------------
    # NON-NEGATIVE METRIC RULES
    # --------------------------------------------------------

    non_negative_metrics = [
        "Revenue",
        "Orders",
        "Customers",
        "Quantity",
        "MarketingSpend",
        "StockLevel",
        "Returns",
    ]

    for metric in non_negative_metrics:
        assert all(
            row[metric] >= 0
            for row in rows
        ), f"Negative {metric} detected in clean dataset."

    # --------------------------------------------------------
    # CROSS-METRIC RULES
    # --------------------------------------------------------

    assert all(
        row["Returns"] <= row["Quantity"]
        for row in rows
    ), "Returns must not exceed Quantity."

    assert all(
        0.0 <= row["Discount"] <= 0.50
        for row in rows
    ), "Discount must remain between 0.00 and 0.50."

    # --------------------------------------------------------
    # NULL VALIDATION
    # --------------------------------------------------------

    assert all(
        all(value is not None for value in row.values())
        for row in rows
    ), "Unexpected NULL detected in clean dataset."

    # --------------------------------------------------------
    # UNIQUE GRAIN VALIDATION
    # --------------------------------------------------------

    grain_keys = [
        (
            row["Date"],
            row["Region"],
            row["Product"],
            row["Channel"],
        )
        for row in rows
    ]

    assert len(grain_keys) == len(set(grain_keys)), (
        "Duplicate Date × Region × Product × Channel grain detected."
    )

def validate_derived_metric_golden_truth(
    rows: list[dict],
) -> None:
    """
    Validates deterministic derived metrics used by the
    Golden Truth specification.
    """

    tolerance = 0.05

    def monthly_total(
        month: str,
        metric: str,
    ) -> float:

        return sum(
            row[metric]
            for row in rows
            if row["Date"].startswith(month)
        )

    may_revenue = monthly_total(
        "2025-05",
        "Revenue",
    )

    june_revenue = monthly_total(
        "2025-06",
        "Revenue",
    )

    may_orders = monthly_total(
        "2025-05",
        "Orders",
    )

    june_orders = monthly_total(
        "2025-06",
        "Orders",
    )

    may_customers = monthly_total(
        "2025-05",
        "Customers",
    )

    june_customers = monthly_total(
        "2025-06",
        "Customers",
    )

    # --------------------------------------------------------
    # AVERAGE ORDER VALUE
    # --------------------------------------------------------

    may_aov = may_revenue / may_orders
    june_aov = june_revenue / june_orders

    aov_change = (
        (june_aov - may_aov)
        / may_aov
        * 100
    )

    assert abs(may_aov - 100.0) <= tolerance

    assert abs(
        june_aov - 152.941176
    ) <= tolerance

    assert abs(
        aov_change - 52.941176
    ) <= tolerance

    # --------------------------------------------------------
    # REVENUE PER CUSTOMER
    # --------------------------------------------------------

    may_rpc = (
        may_revenue
        / may_customers
    )

    june_rpc = (
        june_revenue
        / june_customers
    )

    rpc_change = (
        (june_rpc - may_rpc)
        / may_rpc
        * 100
    )

    assert abs(
        may_rpc - 125.0
    ) <= tolerance

    assert abs(
        june_rpc - 184.659091
    ) <= tolerance

    assert abs(
        rpc_change - 47.727273
    ) <= tolerance

    # --------------------------------------------------------
    # MARMARA CONTRIBUTION
    # --------------------------------------------------------

    may_marmara = sum(
        row["Revenue"]
        for row in rows
        if row["Date"].startswith("2025-05")
        and row["Region"] == "Marmara"
    )

    june_marmara = sum(
        row["Revenue"]
        for row in rows
        if row["Date"].startswith("2025-06")
        and row["Region"] == "Marmara"
    )

    total_change = (
        june_revenue
        - may_revenue
    )

    marmara_change = (
        june_marmara
        - may_marmara
    )

    marmara_contribution = (
        marmara_change
        / total_change
    )

    assert abs(
        marmara_contribution - 0.70
    ) <= 0.01, (
        "Marmara contribution must be approximately 70%."
    )

def print_final_clean_validation_summary(
    rows: list[dict],
) -> None:

    may_revenue = sum(
        row["Revenue"]
        for row in rows
        if row["Date"].startswith("2025-05")
    )

    june_revenue = sum(
        row["Revenue"]
        for row in rows
        if row["Date"].startswith("2025-06")
    )

    may_orders = sum(
        row["Orders"]
        for row in rows
        if row["Date"].startswith("2025-05")
    )

    june_orders = sum(
        row["Orders"]
        for row in rows
        if row["Date"].startswith("2025-06")
    )

    may_customers = sum(
        row["Customers"]
        for row in rows
        if row["Date"].startswith("2025-05")
    )

    june_customers = sum(
        row["Customers"]
        for row in rows
        if row["Date"].startswith("2025-06")
    )

    may_aov = may_revenue / may_orders
    june_aov = june_revenue / june_orders

    may_rpc = may_revenue / may_customers
    june_rpc = june_revenue / june_customers

    print()
    print("Final Clean Dataset Validation")
    print("------------------------------")
    print(f"Rows:                 {len(rows):,}")
    print(f"Columns:              {len(EXPECTED_COLUMNS)}")
    print("Unique grain:         PASSED")
    print("Date coverage:        PASSED")
    print("Dimension integrity:  PASSED")
    print("Metric types:         PASSED")
    print("Non-negative rules:   PASSED")
    print("Cross-metric rules:   PASSED")
    print("NULL rules:           PASSED")

    print()
    print(f"May AOV:              {may_aov:.2f}")
    print(f"June AOV:             {june_aov:.2f}")

    print(
        "AOV Change:           "
        f"{((june_aov - may_aov) / may_aov * 100):.2f}%"
    )

    print()
    print(f"May Revenue/Customer: {may_rpc:.2f}")
    print(f"June Revenue/Customer:{june_rpc:.2f}")

    print(
        "RPC Change:           "
        f"{((june_rpc - may_rpc) / may_rpc * 100):.2f}%"
    )

    print()
    print("Final clean dataset validation: PASSED")

# ============================================================
# DIRTY DATASET INJECTION
# ============================================================

def build_dirty_dataset(
    clean_rows: list[dict],
) -> list[dict]:
    """
    Creates the controlled dirty dataset defined by Golden Truth.

    The clean dataset remains untouched.
    """

    dirty_rows = copy.deepcopy(clean_rows)

    # --------------------------------------------------------
    # 1. STOCK LEVEL MISSING VALUES — 50
    # --------------------------------------------------------

    for index in range(0, 50):
        dirty_rows[index]["StockLevel"] = None

    # --------------------------------------------------------
    # 2. REGION NORMALIZATION ISSUES — 12
    # --------------------------------------------------------

    for index in range(100, 112):
        dirty_rows[index]["Region"] = (
            dirty_rows[index]["Region"].upper()
        )

    # --------------------------------------------------------
    # 3. PRODUCT WHITESPACE ISSUES — 5
    # --------------------------------------------------------

    for index in range(200, 205):
        dirty_rows[index]["Product"] = (
            dirty_rows[index]["Product"] + " "
        )

    # --------------------------------------------------------
    # 4. INVALID PRODUCT-CATEGORY MAPPINGS — 3
    # --------------------------------------------------------

    invalid_mapping_count = 0

    for row in dirty_rows:
        if (
            row["Product"] == "Smart Watch"
            and invalid_mapping_count < 3
        ):
            row["Category"] = "Home"
            invalid_mapping_count += 1

    assert invalid_mapping_count == 3

    # --------------------------------------------------------
    # 5. INVALID DISCOUNT VALUES — 2
    # --------------------------------------------------------

    for index in range(300, 302):
        dirty_rows[index]["Discount"] = 0.75

    # --------------------------------------------------------
    # 6. NEGATIVE ORDER VALUES — 2
    # --------------------------------------------------------

    for index in range(400, 402):
        dirty_rows[index]["Orders"] = -4

    # --------------------------------------------------------
    # 7. RETURNS GREATER THAN QUANTITY — 2
    # --------------------------------------------------------

    for index in range(500, 502):
        dirty_rows[index]["Returns"] = (
            dirty_rows[index]["Quantity"] + 5
        )

    # --------------------------------------------------------
    # 8. INVALID DATE VALUES — 2
    # --------------------------------------------------------

    dirty_rows[600]["Date"] = "not-a-date"
    dirty_rows[601]["Date"] = "2025-13-04"

    # --------------------------------------------------------
    # 9. NUMERIC FORMAT ISSUES — 4
    # --------------------------------------------------------

    for index in range(700, 704):
        original_value = dirty_rows[index]["MarketingSpend"]

        dirty_rows[index]["MarketingSpend"] = (
            f"{original_value:.2f}".replace(".", ",")
    )

    # --------------------------------------------------------
    # 10. DUPLICATE ROWS — 10
    # --------------------------------------------------------

    duplicate_source_rows = dirty_rows[1000:1010]

    dirty_rows.extend(
        copy.deepcopy(duplicate_source_rows)
    )

    return dirty_rows

def validate_clean_dataset_unchanged(
    clean_rows: list[dict],
) -> None:
    """
    Re-validates the clean dataset after dirty dataset creation.

    This proves that dirty injection did not mutate clean rows.
    """

    validate_clean_dataset(clean_rows)
    validate_derived_metric_golden_truth(clean_rows)

def is_region_case_issue(region: str) -> bool:
    """
    Detects case-only Region inconsistencies without treating
    whitespace problems as case problems.
    """

    if region != region.strip():
        return False

    if region in REGIONS:
        return False

    normalized = region.casefold()

    return any(
        normalized == valid_region.casefold()
        for valid_region in REGIONS
    )

def validate_dirty_dataset(
    dirty_rows: list[dict],
) -> None:
    """
    Validates controlled dirty-data issue counts against the
    Golden Truth specification.
    """

    expected_clean_row_count = (
        365
        * len(REGIONS)
        * len(PRODUCT_CATEGORY_MAP)
        * len(CHANNELS)
    )

    expected_dirty_row_count = (
        expected_clean_row_count
        + DIRTY_ISSUE_COUNTS["duplicate_rows"]
    )

    assert len(dirty_rows) == expected_dirty_row_count

    # --------------------------------------------------------
    # STOCK LEVEL NULL
    # --------------------------------------------------------

    stock_nulls = sum(
        1
        for row in dirty_rows
        if row["StockLevel"] is None
    )

    assert stock_nulls == 50

    # --------------------------------------------------------
    # REGION NORMALIZATION
    # --------------------------------------------------------

    region_issues = sum(
        1
        for row in dirty_rows
        if row["Region"] not in REGIONS
        and isinstance(row["Region"], str)
        and any(
            row["Region"].casefold()
            == valid_region.casefold()
            for valid_region in REGIONS
        )
    )

    assert region_issues == 12

    # --------------------------------------------------------
    # PRODUCT WHITESPACE
    # --------------------------------------------------------

    product_whitespace = sum(
        1
        for row in dirty_rows
        if isinstance(row["Product"], str)
        and row["Product"] != row["Product"].strip()
    )

    assert product_whitespace == 5

    # --------------------------------------------------------
    # INVALID PRODUCT-CATEGORY MAPPING
    # --------------------------------------------------------

    invalid_mapping = sum(
        1
        for row in dirty_rows
        if row["Product"] in PRODUCT_CATEGORY_MAP
        and row["Category"]
        != PRODUCT_CATEGORY_MAP[row["Product"]]
    )

    assert invalid_mapping == 3

    # --------------------------------------------------------
    # INVALID DISCOUNT
    # --------------------------------------------------------

    invalid_discount = sum(
        1
        for row in dirty_rows
        if isinstance(row["Discount"], (int, float))
        and not 0.0 <= row["Discount"] <= 0.50
    )

    assert invalid_discount == 2

    # --------------------------------------------------------
    # NEGATIVE ORDERS
    # --------------------------------------------------------

    negative_orders = sum(
        1
        for row in dirty_rows
        if row["Orders"] < 0
    )

    assert negative_orders == 2

    # --------------------------------------------------------
    # RETURNS > QUANTITY
    # --------------------------------------------------------

    invalid_returns = sum(
        1
        for row in dirty_rows
        if row["Returns"] > row["Quantity"]
    )

    assert invalid_returns == 2

    # --------------------------------------------------------
    # INVALID DATES
    # --------------------------------------------------------

    invalid_dates = sum(
        1
        for row in dirty_rows
        if row["Date"] in {
            "not-a-date",
            "2025-13-04",
        }
    )

    assert invalid_dates == 2

    # --------------------------------------------------------
    # NUMERIC FORMAT ISSUES
    # --------------------------------------------------------

    numeric_format_issues = sum(
        1
        for row in dirty_rows
        if isinstance(row["MarketingSpend"], str)
    )

    assert numeric_format_issues == 4

    # --------------------------------------------------------
    # DUPLICATES
    # --------------------------------------------------------

    grain_keys = [
        (
            row["Date"],
            row["Region"],
            row["Product"],
            row["Channel"],
        )
        for row in dirty_rows
    ]

    duplicate_count = (
        len(grain_keys)
        - len(set(grain_keys))
    )

    assert duplicate_count == 10

def print_dirty_dataset_summary(
    dirty_rows: list[dict],
) -> None:

    grain_keys = [
        (
            row["Date"],
            row["Region"],
            row["Product"],
            row["Channel"],
        )
        for row in dirty_rows
    ]

    duplicate_count = (
        len(grain_keys)
        - len(set(grain_keys))
    )

    print()
    print("Dirty Dataset Validation")
    print("------------------------")

    print(f"Dirty rows:                     {len(dirty_rows):,}")
    print("StockLevel missing:             50")
    print("Region normalization issues:    12")
    print("Product whitespace issues:       5")
    print("Invalid Product-Category:         3")
    print("Invalid Discount:                 2")
    print("Negative Orders:                  2")
    print("Returns > Quantity:               2")
    print("Invalid Date:                     2")
    print("Numeric format issues:            4")
    print(
        f"Duplicate grain rows:            {duplicate_count}"
    )

    print()
    print("Dirty dataset validation: PASSED")

# ============================================================
# DATASET EXPORT
# ============================================================

def export_csv(
    rows: list[dict],
    output_path: Path,
) -> None:
    """
    Writes dataset rows to a UTF-8 CSV file using the exact
    Data Dictionary column order.
    """

    with output_path.open(
        mode="w",
        newline="",
        encoding="utf-8-sig",
    ) as csv_file:

        writer = csv.DictWriter(
            csv_file,
            fieldnames=EXPECTED_COLUMNS,
        )

        writer.writeheader()

        for row in rows:
            writer.writerow(
                {
                    column: row[column]
                    for column in EXPECTED_COLUMNS
                }
            )

def export_excel(
    clean_rows: list[dict],
    dirty_rows: list[dict],
    output_path: Path,
) -> None:
    """
    Writes clean and dirty datasets into a single Excel workbook.

    Sheets:
    - CleanData
    - DirtyData
    """

    workbook = Workbook(write_only=True)

    clean_sheet = workbook.create_sheet(
        title="CleanData"
    )

    dirty_sheet = workbook.create_sheet(
        title="DirtyData"
    )

    clean_sheet.append(EXPECTED_COLUMNS)

    for row in clean_rows:
        clean_sheet.append(
            [
                row[column]
                for column in EXPECTED_COLUMNS
            ]
        )

    dirty_sheet.append(EXPECTED_COLUMNS)

    for row in dirty_rows:
        dirty_sheet.append(
            [
                row[column]
                for column in EXPECTED_COLUMNS
            ]
        )

    workbook.save(output_path)

def export_datasets(
    clean_rows: list[dict],
    dirty_rows: list[dict],
) -> None:
    """
    Exports all Day 1 demo dataset artifacts.
    """

    SAMPLE_DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    export_csv(
        rows=clean_rows,
        output_path=CLEAN_CSV_PATH,
    )

    export_csv(
        rows=dirty_rows,
        output_path=DIRTY_CSV_PATH,
    )

    export_excel(
        clean_rows=clean_rows,
        dirty_rows=dirty_rows,
        output_path=EXCEL_PATH,
    )

def validate_exported_files() -> None:
    """
    Validates that all expected exported dataset files exist
    and are non-empty.
    """

    expected_files = [
        CLEAN_CSV_PATH,
        DIRTY_CSV_PATH,
        EXCEL_PATH,
    ]

    for file_path in expected_files:

        assert file_path.exists(), (
            f"Expected export file does not exist: "
            f"{file_path}"
        )

        assert file_path.stat().st_size > 0, (
            f"Export file is empty: "
            f"{file_path}"
        )

def validate_exported_csv_files() -> None:
    """
    Reads exported CSV files back from disk and validates
    their structure and row counts.
    """

    with CLEAN_CSV_PATH.open(
        mode="r",
        newline="",
        encoding="utf-8-sig",
    ) as clean_file:

        clean_reader = csv.DictReader(clean_file)

        assert clean_reader.fieldnames == EXPECTED_COLUMNS, (
            "Clean CSV column order mismatch."
        )

        clean_row_count = sum(
            1
            for _ in clean_reader
        )

    assert clean_row_count == 43_800, (
        f"Clean CSV row count mismatch: "
        f"{clean_row_count} != 43800"
    )

    with DIRTY_CSV_PATH.open(
        mode="r",
        newline="",
        encoding="utf-8-sig",
    ) as dirty_file:

        dirty_reader = csv.DictReader(dirty_file)

        assert dirty_reader.fieldnames == EXPECTED_COLUMNS, (
            "Dirty CSV column order mismatch."
        )

        dirty_row_count = sum(
            1
            for _ in dirty_reader
        )

    expected_dirty_row_count = (
        43_800
        + DIRTY_ISSUE_COUNTS["duplicate_rows"]
    )

    assert dirty_row_count == expected_dirty_row_count, (
        f"Dirty CSV row count mismatch: "
        f"{dirty_row_count} != {expected_dirty_row_count}"
    )
def validate_exported_excel_file() -> None:
    """
    Opens the generated Excel workbook and validates
    sheet names, headers, row counts, and column counts.

    Rows are counted explicitly because max_row may not be
    available reliably in read-only mode.
    """

    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True,
    )

    try:
        assert workbook.sheetnames == [
            "CleanData",
            "DirtyData",
        ], (
            f"Unexpected Excel sheets: "
            f"{workbook.sheetnames}"
        )

        clean_sheet = workbook["CleanData"]
        dirty_sheet = workbook["DirtyData"]

        # ----------------------------------------------------
        # CLEAN SHEET
        # ----------------------------------------------------

        clean_rows = clean_sheet.iter_rows(
            values_only=True
        )

        clean_header = next(clean_rows)

        assert list(clean_header) == EXPECTED_COLUMNS, (
            "CleanData Excel header mismatch."
        )

        clean_data_row_count = sum(
            1
            for _ in clean_rows
        )

        assert clean_data_row_count == 43_800, (
            f"CleanData Excel data row count mismatch: "
            f"{clean_data_row_count} != 43800"
        )

        # ----------------------------------------------------
        # DIRTY SHEET
        # ----------------------------------------------------

        dirty_rows = dirty_sheet.iter_rows(
            values_only=True
        )

        dirty_header = next(dirty_rows)

        assert list(dirty_header) == EXPECTED_COLUMNS, (
            "DirtyData Excel header mismatch."
        )

        dirty_data_row_count = sum(
            1
            for _ in dirty_rows
        )

        expected_dirty_row_count = (
            43_800
            + DIRTY_ISSUE_COUNTS["duplicate_rows"]
        )

        assert dirty_data_row_count == expected_dirty_row_count, (
            f"DirtyData Excel data row count mismatch: "
            f"{dirty_data_row_count} "
            f"!= {expected_dirty_row_count}"
        )

        # ----------------------------------------------------
        # COLUMN COUNTS
        # ----------------------------------------------------

        assert len(clean_header) == len(EXPECTED_COLUMNS), (
            "CleanData Excel column count mismatch."
        )

        assert len(dirty_header) == len(EXPECTED_COLUMNS), (
            "DirtyData Excel column count mismatch."
        )

    finally:
        workbook.close()

def print_export_summary() -> None:

    print()
    print("Dataset Export")
    print("--------------")

    print(
        f"Clean CSV: {CLEAN_CSV_PATH.name}"
    )

    print(
        f"Dirty CSV: {DIRTY_CSV_PATH.name}"
    )

    print(
        f"Excel:     {EXCEL_PATH.name}"
    )

    print()
    print("Exported file validation: PASSED")

def validate_business_relevance_scenario(
    rows: list[dict],
) -> None:
    """
    Validates the August-to-September controlled Business
    Relevance scenario.
    """

    tolerance = 0.0001

    def monthly_average(
        month: str,
        metric: str,
    ) -> float:
        month_values = [
            row[metric]
            for row in rows
            if row["Date"].startswith(month)
        ]

        return sum(month_values) / len(month_values)

    def monthly_return_rate(month: str) -> float:
        total_returns = sum(
            row["Returns"]
            for row in rows
            if row["Date"].startswith(month)
        )

        total_quantity = sum(
            row["Quantity"]
            for row in rows
            if row["Date"].startswith(month)
        )

        return total_returns / total_quantity

    august_discount = monthly_average(
        "2025-08",
        "Discount",
    )

    september_discount = monthly_average(
        "2025-09",
        "Discount",
    )

    august_return_rate = monthly_return_rate(
        "2025-08"
    )

    september_return_rate = monthly_return_rate(
        "2025-09"
    )

    assert abs(
        august_discount - 0.10
    ) <= tolerance

    assert abs(
        september_discount - 0.18
    ) <= tolerance

    assert abs(
        august_return_rate - 0.020
    ) <= tolerance

    assert abs(
        september_return_rate - 0.035
    ) <= tolerance

def print_remaining_fields_summary(
    rows: list[dict],
) -> None:

    def monthly_quantity(month: str) -> float:
        return sum(
            row["Quantity"]
            for row in rows
            if row["Date"].startswith(month)
        )

    def monthly_return_rate(month: str) -> float:
        total_returns = sum(
            row["Returns"]
            for row in rows
            if row["Date"].startswith(month)
        )

        total_quantity = monthly_quantity(month)

        return total_returns / total_quantity

    def monthly_discount(month: str) -> float:
        values = [
            row["Discount"]
            for row in rows
            if row["Date"].startswith(month)
        ]

        return sum(values) / len(values)

    print()
    print("Remaining Fields Validation")
    print("---------------------------")

    print(
        f"August Discount:       "
        f"{monthly_discount('2025-08') * 100:.2f}%"
    )

    print(
        f"September Discount:    "
        f"{monthly_discount('2025-09') * 100:.2f}%"
    )

    print(
        f"August ReturnRate:     "
        f"{monthly_return_rate('2025-08') * 100:.2f}%"
    )

    print(
        f"September ReturnRate:  "
        f"{monthly_return_rate('2025-09') * 100:.2f}%"
    )

    print()
    print("Remaining fields validation: PASSED")

def validate_controlled_metrics(
    rows: list[dict],
) -> None:
    """
    Validates monthly Orders, Customers, and MarketingSpend
    against the controlled Golden Scenario targets.
    """

    tolerance = 0.01

    target_sets = {
        "Orders": MONTHLY_ORDER_TARGETS,
        "Customers": MONTHLY_CUSTOMER_TARGETS,
        "MarketingSpend": MONTHLY_MARKETING_SPEND_TARGETS,
    }

    for metric_name, targets in target_sets.items():

        for month, expected_total in targets.items():

            actual_total = sum(
                row[metric_name]
                for row in rows
                if row["Date"].startswith(month)
            )

            assert abs(actual_total - expected_total) <= tolerance, (
                f"{metric_name} mismatch for {month}: "
                f"{actual_total} != {expected_total}"
            )
def validate_main_change_scenario(
    rows: list[dict],
) -> None:
    """
    Validates the main May-to-June Golden Scenario changes.
    """

    tolerance = 0.01

    def monthly_total(month: str, metric: str) -> float:
        return sum(
            row[metric]
            for row in rows
            if row["Date"].startswith(month)
        )

    may_orders = monthly_total("2025-05", "Orders")
    june_orders = monthly_total("2025-06", "Orders")

    may_customers = monthly_total("2025-05", "Customers")
    june_customers = monthly_total("2025-06", "Customers")

    may_marketing = monthly_total("2025-05", "MarketingSpend")
    june_marketing = monthly_total("2025-06", "MarketingSpend")

    orders_change = (
        (june_orders - may_orders)
        / may_orders
        * 100
    )

    customers_change = (
        (june_customers - may_customers)
        / may_customers
        * 100
    )

    marketing_change = (
        (june_marketing - may_marketing)
        / may_marketing
        * 100
    )

    assert abs(orders_change - (-15.0)) <= tolerance, (
        f"Orders Golden Change mismatch: {orders_change}"
    )

    assert abs(customers_change - (-12.0)) <= tolerance, (
        f"Customers Golden Change mismatch: {customers_change}"
    )

    assert abs(marketing_change - 20.0) <= tolerance, (
        f"MarketingSpend Golden Change mismatch: {marketing_change}"
    )
def print_controlled_metrics_summary(
    rows: list[dict],
) -> None:

    def monthly_total(month: str, metric: str) -> float:
        return sum(
            row[metric]
            for row in rows
            if row["Date"].startswith(month)
        )

    print()
    print("Controlled Metrics Validation")
    print("-----------------------------")

    for month in ["2025-05", "2025-06"]:
        print()
        print(month)
        print(
            f"  Orders:         "
            f"{monthly_total(month, 'Orders'):,.2f}"
        )
        print(
            f"  Customers:      "
            f"{monthly_total(month, 'Customers'):,.2f}"
        )
        print(
            f"  MarketingSpend: "
            f"{monthly_total(month, 'MarketingSpend'):,.2f}"
        )

    may_orders = monthly_total("2025-05", "Orders")
    june_orders = monthly_total("2025-06", "Orders")

    may_customers = monthly_total("2025-05", "Customers")
    june_customers = monthly_total("2025-06", "Customers")

    may_marketing = monthly_total("2025-05", "MarketingSpend")
    june_marketing = monthly_total("2025-06", "MarketingSpend")

    print()
    print(
        "Orders Change: "
        f"{((june_orders - may_orders) / may_orders * 100):.2f}%"
    )

    print(
        "Customers Change: "
        f"{((june_customers - may_customers) / may_customers * 100):.2f}%"
    )

    print(
        "MarketingSpend Change: "
        f"{((june_marketing - may_marketing) / may_marketing * 100):.2f}%"
    )

    print()
    print("Controlled metrics validation: PASSED")

def validate_daily_revenue_rows(
    rows: list[dict],
) -> None:
    """
    Ensures that converting monthly Revenue targets to daily rows
    does not change Golden Scenario totals.
    """

    tolerance = 0.01

    expected_row_count = (
        365
        * len(REGIONS)
        * len(PRODUCT_CATEGORY_MAP)
        * len(CHANNELS)
    )

    assert len(rows) == expected_row_count, (
        f"Daily row count mismatch: "
        f"{len(rows)} != {expected_row_count}"
    )

    # --------------------------------------------------------
    # MONTHLY TOTAL REVENUE
    # --------------------------------------------------------

    for month, expected_total in MONTHLY_REVENUE_TARGETS.items():

        actual_total = sum(
            row["Revenue"]
            for row in rows
            if row["Date"].startswith(month)
        )

        assert abs(actual_total - expected_total) <= tolerance, (
            f"Daily Revenue aggregation mismatch for {month}: "
            f"{actual_total} != {expected_total}"
        )

    # --------------------------------------------------------
    # MAY / JUNE REGION TARGETS
    # --------------------------------------------------------

    for month in ["2025-05", "2025-06"]:

        for region, expected_total in REGION_REVENUE_TARGETS[month].items():

            actual_total = sum(
                row["Revenue"]
                for row in rows
                if row["Date"].startswith(month)
                and row["Region"] == region
            )

            assert abs(actual_total - expected_total) <= tolerance, (
                f"Daily Region Revenue mismatch for "
                f"{month} / {region}: "
                f"{actual_total} != {expected_total}"
            )

    # --------------------------------------------------------
    # MAY / JUNE CHANNEL TARGETS
    # --------------------------------------------------------

    for month in ["2025-05", "2025-06"]:

        for channel, expected_total in CHANNEL_REVENUE_TARGETS[month].items():

            actual_total = sum(
                row["Revenue"]
                for row in rows
                if row["Date"].startswith(month)
                and row["Channel"] == channel
            )

            assert abs(actual_total - expected_total) <= tolerance, (
                f"Daily Channel Revenue mismatch for "
                f"{month} / {channel}: "
                f"{actual_total} != {expected_total}"
            )

    # --------------------------------------------------------
    # CONTROLLED PRODUCT TARGETS
    # --------------------------------------------------------

    controlled_products = {
        "Smart Watch": SMART_WATCH_MONTHLY_REVENUE,
        "Office Chair": OFFICE_CHAIR_MONTHLY_REVENUE,
        "Bluetooth Speaker": BLUETOOTH_SPEAKER_MONTHLY_REVENUE,
    }

    for product, targets in controlled_products.items():

        for month, expected_total in targets.items():

            actual_total = sum(
                row["Revenue"]
                for row in rows
                if row["Date"].startswith(month)
                and row["Product"] == product
            )

            assert abs(actual_total - expected_total) <= tolerance, (
                f"Daily Product Revenue mismatch for "
                f"{month} / {product}: "
                f"{actual_total} != {expected_total}"
            )
def print_daily_distribution_summary(
    rows: list[dict],
) -> None:

    print()
    print("Daily Distribution Validation")
    print("-----------------------------")
    print(f"Generated daily rows: {len(rows):,}")

    june_total = sum(
        row["Revenue"]
        for row in rows
        if row["Date"].startswith("2025-06")
    )

    june_marmara = sum(
        row["Revenue"]
        for row in rows
        if row["Date"].startswith("2025-06")
        and row["Region"] == "Marmara"
    )

    june_smart_watch = sum(
        row["Revenue"]
        for row in rows
        if row["Date"].startswith("2025-06")
        and row["Product"] == "Smart Watch"
    )

    june_online = sum(
        row["Revenue"]
        for row in rows
        if row["Date"].startswith("2025-06")
        and row["Channel"] == "Online"
    )

    print(f"June Total Revenue:       {june_total:,.2f}")
    print(f"June Marmara Revenue:     {june_marmara:,.2f}")
    print(f"June Smart Watch Revenue: {june_smart_watch:,.2f}")
    print(f"June Online Revenue:      {june_online:,.2f}")

    print()
    print("Daily distribution validation: PASSED")

def print_revenue_validation_summary(
    revenue_cubes: dict[str, dict[tuple[str, str, str], float]],
) -> None:

    print()
    print("Revenue Allocation Validation")
    print("-----------------------------")

    for month in ["2025-05", "2025-06"]:

        cube = revenue_cubes[month]

        total = sum(cube.values())

        marmara = sum(
            revenue
            for (region, product, channel), revenue in cube.items()
            if region == "Marmara"
        )

        smart_watch = sum(
            revenue
            for (region, product, channel), revenue in cube.items()
            if product == "Smart Watch"
        )

        online = sum(
            revenue
            for (region, product, channel), revenue in cube.items()
            if channel == "Online"
        )

        print()
        print(month)
        print(f"  Total Revenue:      {total:,.2f}")
        print(f"  Marmara Revenue:    {marmara:,.2f}")
        print(f"  Smart Watch Revenue:{smart_watch:,.2f}")
        print(f"  Online Revenue:     {online:,.2f}")

    print()
    print("Revenue allocation validation: PASSED")
# ============================================================
# BASIC CONFIG VALIDATION
# ============================================================

def validate_configuration() -> None:
    """
    Validates static generator configuration before any dataset
    generation takes place.

    This catches specification mistakes early.
    """

    assert len(REGIONS) == 5, "Expected exactly 5 regions."
    assert len(PRODUCT_CATEGORY_MAP) == 12, "Expected exactly 12 products."
    assert len(CHANNELS) == 2, "Expected exactly 2 sales channels."

    assert set(PRODUCT_CATEGORY_MAP.values()) == {
        "Electronics",
        "Home",
        "Office",
    }, "Unexpected product categories."

    for month, total_revenue in MONTHLY_REVENUE_TARGETS.items():
        assert total_revenue > 0, (
            f"Monthly Revenue must be positive: {month}"
        )

    for month in ["2025-05", "2025-06"]:
        region_total = sum(REGION_REVENUE_TARGETS[month].values())
        expected_total = MONTHLY_REVENUE_TARGETS[month]

        assert region_total == expected_total, (
            f"Region Revenue total mismatch for {month}: "
            f"{region_total} != {expected_total}"
        )

        channel_total = sum(CHANNEL_REVENUE_TARGETS[month].values())

        assert channel_total == expected_total, (
            f"Channel Revenue total mismatch for {month}: "
            f"{channel_total} != {expected_total}"
        )

    assert MONTHLY_ORDER_TARGETS["2025-05"] == 10_000
    assert MONTHLY_ORDER_TARGETS["2025-06"] == 8_500

    assert MONTHLY_CUSTOMER_TARGETS["2025-05"] == 8_000
    assert MONTHLY_CUSTOMER_TARGETS["2025-06"] == 7_040

    assert MONTHLY_MARKETING_SPEND_TARGETS["2025-05"] == 100_000
    assert MONTHLY_MARKETING_SPEND_TARGETS["2025-06"] == 120_000


def print_configuration_summary() -> None:
    print("Data Story Builder — Demo Dataset Generator")
    print("-------------------------------------------")
    print(f"Period: {CONFIG.start_date} -> {CONFIG.end_date}")
    print(f"Frequency: {CONFIG.frequency}")
    print(f"Random seed: {CONFIG.random_seed}")
    print(f"Regions: {len(REGIONS)}")
    print(f"Products: {len(PRODUCT_CATEGORY_MAP)}")
    print(f"Channels: {len(CHANNELS)}")

    expected_rows = (
        365
        * len(REGIONS)
        * len(PRODUCT_CATEGORY_MAP)
        * len(CHANNELS)
    )

    print(f"Expected clean row count: {expected_rows:,}")
    print()
    print("Configuration validation: PASSED")


def main() -> None:
    SAMPLE_DATA_DIR.mkdir(parents=True, exist_ok=True)

    initialize_random_seed()

    validate_configuration()
    print_configuration_summary()

    revenue_cubes = build_and_validate_revenue_cubes()
    print_revenue_validation_summary(revenue_cubes)

    daily_rows = build_daily_revenue_rows(revenue_cubes)

    validate_daily_revenue_rows(daily_rows)
    print_daily_distribution_summary(daily_rows)

    add_controlled_metrics(daily_rows)

    validate_controlled_metrics(daily_rows)
    validate_main_change_scenario(daily_rows)

    print_controlled_metrics_summary(daily_rows)

    add_remaining_dataset_fields(daily_rows)

    validate_remaining_dataset_fields(daily_rows)
    validate_business_relevance_scenario(daily_rows)

    print_remaining_fields_summary(daily_rows)

    validate_clean_dataset(daily_rows)
    validate_derived_metric_golden_truth(daily_rows)

    print_final_clean_validation_summary(daily_rows)

    dirty_rows = build_dirty_dataset(daily_rows)

    validate_clean_dataset_unchanged(daily_rows)
    validate_dirty_dataset(dirty_rows)

    print_dirty_dataset_summary(dirty_rows)

    export_datasets(
        clean_rows=daily_rows,
        dirty_rows=dirty_rows,
    )

    validate_exported_files()
    validate_exported_csv_files()
    validate_exported_excel_file()

    print_export_summary()
    
if __name__ == "__main__":
    main()